"""Tests for TXT/DOCX/EPUB converters: content fidelity and explicit RTL
formatting artifacts (OXML bidi flags for DOCX, dir="rtl"/CSS for EPUB).
"""
from __future__ import annotations

import zipfile

from docx import Document
from ebooklib import ITEM_DOCUMENT, epub

from common.models import Book, PageText
from converters.docx_writer import write_docx
from converters.epub_writer import write_epub
from converters.fonts import DEFAULT_PERSIAN_FONT, get_font_file_path, get_persian_font_name
from converters.txt import write_txt


def _sample_book() -> Book:
    return Book(
        title="کتاب آزمایشی",
        author="نویسنده آزمایشی",
        pages=[
            PageText(page_number=2, text="این متن صفحه دوم است."),
            PageText(page_number=1, text="این متن صفحه اول است."),
        ],
    )


def test_get_persian_font_name_defaults_and_override(monkeypatch):
    monkeypatch.delenv("PERSIAN_FONT_NAME", raising=False)
    assert get_persian_font_name() == DEFAULT_PERSIAN_FONT
    assert get_persian_font_name("CustomFont") == "CustomFont"

    monkeypatch.setenv("PERSIAN_FONT_NAME", "EnvFont")
    assert get_persian_font_name() == "EnvFont"


def test_get_font_file_path_defaults_to_none(monkeypatch):
    monkeypatch.delenv("PERSIAN_FONT_PATH", raising=False)
    assert get_font_file_path() is None
    assert get_font_file_path("/custom/font.ttf") == "/custom/font.ttf"


def test_write_txt_orders_pages_and_includes_title(tmp_path):
    book = _sample_book()
    output_path = write_txt(book, tmp_path / "book.txt")

    assert output_path.exists()
    content = output_path.read_text(encoding="utf-8")

    assert book.title in content
    # page 1 text must appear before page 2 text despite input order
    assert content.index("صفحه اول") < content.index("صفحه دوم")


def test_write_docx_produces_rtl_flags_and_readable_paragraphs(tmp_path):
    book = _sample_book()
    output_path = write_docx(book, tmp_path / "book.docx")

    assert output_path.exists()

    document = Document(str(output_path))
    assert len(document.paragraphs) >= 2

    # The document should default to a right-to-left section.
    section_xml = document.sections[0]._sectPr.xml
    assert "w:bidi" in section_xml

    # At least one paragraph/run must carry explicit bidi/rtl OXML markers.
    combined_xml = "".join(p._p.xml for p in document.paragraphs)
    assert "w:bidi" in combined_xml
    assert "w:rtl" in combined_xml

    body_text = "\n".join(p.text for p in document.paragraphs)
    assert "صفحه اول" in body_text
    assert "صفحه دوم" in body_text


def test_write_epub_produces_valid_rtl_epub(tmp_path):
    book = _sample_book()
    output_path = write_epub(book, tmp_path / "book.epub")

    assert output_path.exists()

    # A well-formed EPUB is a valid zip archive.
    assert zipfile.is_zipfile(output_path)

    read_back = epub.read_epub(str(output_path))
    assert read_back.get_metadata("DC", "title")[0][0] == book.title

    html_items = [
        item
        for item in read_back.get_items()
        if item.get_type() == ITEM_DOCUMENT and not isinstance(item, epub.EpubNav)
    ]
    assert len(html_items) == 2

    # Inspect the actual persisted XHTML page files inside the EPUB (a zip
    # archive) directly, rather than re-reading them through ebooklib (which
    # would regenerate content from a freshly-parsed EpubHtml object that
    # does not restore the per-chapter `lang`/`direction` Python attributes,
    # masking whether RTL markup was actually written to disk).
    with zipfile.ZipFile(output_path) as archive:
        page_files = [
            name
            for name in archive.namelist()
            if name.endswith(".xhtml") and "page_" in name
        ]
        assert len(page_files) == 2
        all_content = "".join(archive.read(name).decode("utf-8") for name in page_files)

    assert 'dir="rtl"' in all_content
    assert "صفحه اول" in all_content
    assert "صفحه دوم" in all_content
